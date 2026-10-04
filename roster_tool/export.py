"""Export a roster to CSV or Excel.

This file is commented line by line. CSV is a plain text table any spreadsheet
can open; Excel export (via the optional ``openpyxl`` package) adds colours,
frozen headers and a legend.
"""

from __future__ import annotations

# The standard-library CSV writer; ``re`` strips characters Excel rejects.
import csv
import re

# Codes, descriptions and weekday names used to build the table.
from .model import ALL_CODES, CODE_DESCRIPTIONS, WEEKDAY_NAMES, Roster
# Issue is only imported for the type hint on to_excel.
from .validator import Issue

# Control characters that openpyxl refuses to put in a worksheet (they can
# sneak in when names/notes are pasted from other apps). We strip them so an
# Excel export never fails with IllegalCharacterError.
_ILLEGAL_XLSX = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _clean(value):
    # Remove illegal control characters from a string; leave non-strings as-is.
    return _ILLEGAL_XLSX.sub("", value) if isinstance(value, str) else value

# Cell background colours (hex, no '#') shared by the GUI grid and Excel export.
CODE_COLORS = {
    "M": "FEF3C7",
    "E": "DBEAFE",
    "N": "312E81",
    "CO": "D1FAE5",
    "L": "FEE2E2",
    "LL": "FCA5A5",
    "WO": "F3F4F6",
    "H": "EDE9FE",
}
# Matching text colours so each code stays readable on its background.
CODE_TEXT_COLORS = {
    "M": "92400E",
    "E": "1E40AF",
    "N": "FFFFFF",
    "CO": "065F46",
    "L": "991B1B",
    "LL": "7F1D1D",
    "WO": "9CA3AF",
    "H": "5B21B6",
}
# Header shading per day type.
DAYTYPE_COLORS = {"Working": "FFFFFF", "Freeze": "FFEDD5", "Weekend": "F3F4F6", "Holiday": "EDE9FE"}

# The summary columns appended to the right of the grid.
SUMMARY_KEYS = ("M", "E", "N", "CO", "L", "LL", "Primary", "Secondary")


def table_rows(roster: Roster) -> list[list[str]]:
    """The roster as a plain table: header rows, one row per engineer, then
    the on-call rows."""
    cfg = roster.config
    days = roster.days
    # Row 1: column headers - the fixed columns, each day number, then summaries.
    header = ["Engineer", "Designation", "SME"] + [d.strftime("%d") for d in days] + list(SUMMARY_KEYS)
    # Row 2: the weekday under each day number.
    weekday = ["", "", ""] + [WEEKDAY_NAMES[d.weekday()] for d in days] + [""] * len(SUMMARY_KEYS)
    # Row 3: the day type (Working/Weekend/Holiday/Freeze).
    daytype = ["", "", "Day type"] + [cfg.day_type(d) for d in days] + [""] * len(SUMMARY_KEYS)
    # The first three rows are the header block.
    rows = [header, weekday, daytype]
    # One row per engineer: their details, their codes, then their tallies.
    for e in cfg.engineers:
        counts = roster.counts(e.name)
        rows.append(
            [e.name, e.designation, "Yes" if e.is_sme else "No"]
            + [roster.code(e.name, d) for d in days]
            + [str(counts[k]) for k in SUMMARY_KEYS]
        )
    # Two on-call rows at the bottom.
    rows.append(["Primary on-call", "", ""] + [roster.primary.get(d, "") for d in days] + [""] * len(SUMMARY_KEYS))
    rows.append(["Secondary on-call", "", ""] + [roster.secondary.get(d, "") for d in days] + [""] * len(SUMMARY_KEYS))
    return rows


def to_csv(roster: Roster, path: str) -> None:
    # Write the plain table plus a small legend to a CSV file.
    with open(path, "w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerows(table_rows(roster))   # the whole grid
        writer.writerow([])                    # a blank separator row
        writer.writerow(["Legend"])            # legend heading
        for code in ALL_CODES:
            writer.writerow([code, CODE_DESCRIPTIONS[code]])   # code + meaning


def to_excel(roster: Roster, path: str, issues: list[Issue] | None = None) -> None:
    # openpyxl gives the nicest output, but it is an optional package. When it
    # is not installed, fall back to our own standard-library .xlsx writer so
    # Excel export works everywhere with nothing to install.
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        _to_xlsx_stdlib(roster, path, issues)
        return

    cfg = roster.config
    # Create the workbook and its first sheet.
    wb = Workbook()
    ws = wb.active
    ws.title = f"Roster {cfg.year}-{cfg.month:02d}"
    # A thin grey border applied to every cell.
    thin = Side(style="thin", color="BBBBBB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    # Centre alignment for the day/summary cells.
    center = Alignment(horizontal="center", vertical="center")

    def fill(hex_color):
        # A solid background fill of the given colour.
        return PatternFill("solid", start_color=hex_color, end_color=hex_color)

    # The same table used for CSV.
    rows = table_rows(roster)
    first_day_col = 4          # columns 1-3 are Engineer/Designation/SME
    n_days = len(roster.days)
    # Write and style every cell (openpyxl rows/cols are 1-based).
    for r, row in enumerate(rows, 1):
        for c, value in enumerate(row, 1):
            # Strip any characters Excel would reject (from pasted names etc.).
            value = _clean(value)
            # Summary numbers become real numbers; everything else stays text.
            cell = ws.cell(row=r, column=c, value=int(value) if value.isdigit() and r > 3 else value)
            cell.border = border
            if c >= first_day_col:
                cell.alignment = center       # centre the day/summary columns
            if r <= 3:
                cell.font = Font(bold=True)    # bold the three header rows
            # Which day (if any) this column corresponds to.
            day_idx = c - first_day_col
            if 0 <= day_idx < n_days:
                d = roster.days[day_idx]
                if r <= 3:
                    # Header cells get their day-type colour.
                    cell.fill = fill(DAYTYPE_COLORS[cfg.day_type(d)])
                elif value in CODE_COLORS:
                    # Body cells get their code colour and matching text colour.
                    cell.fill = fill(CODE_COLORS[value])
                    cell.font = Font(color=CODE_TEXT_COLORS.get(value, "000000"), bold=value == "N")
    # Keep the header rows and name columns visible while scrolling.
    ws.freeze_panes = ws.cell(row=4, column=first_day_col)
    # Fixed widths for the first three columns.
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 9
    # On-call rows hold names, so size each day column to fit them.
    for i, d in enumerate(roster.days):
        width = max(5, len(roster.primary.get(d, "")) + 1, len(roster.secondary.get(d, "")) + 1)
        ws.column_dimensions[get_column_letter(first_day_col + i)].width = width

    # Write the legend a couple of rows below the table.
    legend_row = len(rows) + 2
    ws.cell(row=legend_row, column=1, value="Legend").font = Font(bold=True)
    for i, code in enumerate(ALL_CODES, 1):
        c = ws.cell(row=legend_row + i, column=1, value=code)
        c.fill = fill(CODE_COLORS[code])
        c.font = Font(color=CODE_TEXT_COLORS.get(code, "000000"))
        ws.cell(row=legend_row + i, column=2, value=CODE_DESCRIPTIONS[code])

    # If issues were supplied, add a second "Issues" sheet listing them.
    if issues is not None:
        ws2 = wb.create_sheet("Issues")
        ws2.append(["Severity", "Date", "Message"])
        for i in issues:
            # Issue messages can contain engineer names, so clean them too.
            ws2.append([i.severity, i.day.isoformat() if i.day else "", _clean(i.message)])
        if not issues:
            ws2.append(["OK", "", "All mandatory rules are satisfied"])
        ws2.column_dimensions["C"].width = 90

    # Save the finished workbook.
    wb.save(path)


# ---- Standard-library .xlsx writer ------------------------------------------
# Used when openpyxl is not installed. An .xlsx file is just a ZIP of XML
# parts (the Office Open XML format); we build the minimal set of parts by
# hand so Excel export needs no third-party package.

def _col_letter(n: int) -> str:
    # 1 -> "A", 27 -> "AA", etc. (spreadsheet column letters).
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def _xml_escape(text: str) -> str:
    # Escape the five XML special characters.
    return (text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            .replace('"', "&quot;").replace("'", "&apos;"))


class _Styles:
    """Builds the <fonts>/<fills>/<borders>/<cellXfs> tables for styles.xml,
    de-duplicating as it goes and handing back a style index per cell."""

    def __init__(self):
        # Fill 0 (none) and 1 (gray125) are required to exist by the format.
        self.fonts = ['<font><sz val="11"/><name val="Calibri"/></font>']
        self.fills = ['<fill><patternFill patternType="none"/></fill>',
                      '<fill><patternFill patternType="gray125"/></fill>']
        # Border 0 = none, border 1 = thin grey box.
        self.borders = ['<border><left/><right/><top/><bottom/><diagonal/></border>',
                        '<border><left style="thin"><color rgb="FFBBBBBB"/></left>'
                        '<right style="thin"><color rgb="FFBBBBBB"/></right>'
                        '<top style="thin"><color rgb="FFBBBBBB"/></top>'
                        '<bottom style="thin"><color rgb="FFBBBBBB"/></bottom><diagonal/></border>']
        # cellXf 0 = the default style.
        self.xfs = ['<xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>']

    def _font(self, rgb="000000", bold=False):
        # Return the index of a font with this colour/boldness, adding it if new.
        xml = "<font>" + ("<b/>" if bold else "") + '<sz val="11"/>' + \
            f'<color rgb="FF{rgb}"/>' + '<name val="Calibri"/></font>'
        if xml not in self.fonts:
            self.fonts.append(xml)
        return self.fonts.index(xml)

    def _fill(self, rgb):
        # Return the index of a solid fill of this colour (0 for no fill).
        if not rgb:
            return 0
        xml = f'<fill><patternFill patternType="solid"><fgColor rgb="FF{rgb}"/></patternFill></fill>'
        if xml not in self.fills:
            self.fills.append(xml)
        return self.fills.index(xml)

    def style(self, *, color="000000", bold=False, fill="", center=False, border=True):
        # Return a cellXf index for the given appearance, de-duplicated.
        font_id = self._font(color, bold)
        fill_id = self._fill(fill)
        border_id = 1 if border else 0
        align = '<alignment horizontal="center" vertical="center"/>' if center else ""
        applies = ' applyAlignment="1"' if center else ""
        xf = (f'<xf numFmtId="0" fontId="{font_id}" fillId="{fill_id}" borderId="{border_id}" '
              f'xfId="0" applyFont="1" applyFill="1" applyBorder="1"{applies}>{align}</xf>')
        if xf not in self.xfs:
            self.xfs.append(xf)
        return self.xfs.index(xf)

    def xml(self):
        # Assemble the full styles.xml document.
        return (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            f'<fonts count="{len(self.fonts)}">{"".join(self.fonts)}</fonts>'
            f'<fills count="{len(self.fills)}">{"".join(self.fills)}</fills>'
            f'<borders count="{len(self.borders)}">{"".join(self.borders)}</borders>'
            '<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>'
            f'<cellXfs count="{len(self.xfs)}">{"".join(self.xfs)}</cellXfs>'
            '<cellStyles count="1"><cellStyle name="Normal" xfId="0" builtinId="0"/></cellStyles>'
            '</styleSheet>'
        )


def _cell(ref, value, style_id):
    # One <c> cell: a number, inline string, or empty cell.
    if value is None or value == "":
        return f'<c r="{ref}" s="{style_id}"/>'
    if isinstance(value, int):
        return f'<c r="{ref}" s="{style_id}"><v>{value}</v></c>'
    return f'<c r="{ref}" s="{style_id}" t="inlineStr"><is><t xml:space="preserve">{_xml_escape(str(value))}</t></is></c>'


def _sheet_xml(rows, freeze=None, col_widths=None):
    # Build one worksheet XML from a list of rows, where each row is a list of
    # (value, style_id) pairs. ``freeze`` is (rows, cols); ``col_widths`` maps
    # a 1-based column index to a width.
    out = ['<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
           '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">']
    # Frozen header rows/columns.
    if freeze:
        fr, fc = freeze
        top_left = f"{_col_letter(fc + 1)}{fr + 1}"
        out.append('<sheetViews><sheetView workbookViewId="0">'
                   f'<pane xSplit="{fc}" ySplit="{fr}" topLeftCell="{top_left}" activePane="bottomRight" state="frozen"/>'
                   '</sheetView></sheetViews>')
    # Column widths.
    if col_widths:
        cols = "".join(f'<col min="{c}" max="{c}" width="{w}" customWidth="1"/>' for c, w in sorted(col_widths.items()))
        out.append(f"<cols>{cols}</cols>")
    # The rows and cells.
    out.append("<sheetData>")
    for r, row in enumerate(rows, 1):
        cells = "".join(_cell(f"{_col_letter(c)}{r}", value, style_id) for c, (value, style_id) in enumerate(row, 1))
        out.append(f'<row r="{r}">{cells}</row>')
    out.append("</sheetData></worksheet>")
    return "".join(out)


def _to_xlsx_stdlib(roster: Roster, path: str, issues: list[Issue] | None = None) -> None:
    # Build a styled .xlsx using only the standard library (zipfile + XML).
    import zipfile

    cfg = roster.config
    st = _Styles()
    # Pre-build the styles we reuse.
    name_style = st.style(center=False)                        # plain bordered (names)
    num_style = st.style(center=True)                          # centred number
    header_styles = {dt: st.style(bold=True, fill=DAYTYPE_COLORS[dt], center=True) for dt in DAYTYPE_COLORS}
    header_plain = st.style(bold=True, center=True)
    code_styles = {code: st.style(color=CODE_TEXT_COLORS.get(code, "000000"), bold=(code == "N"),
                                  fill=CODE_COLORS[code], center=True) for code in ALL_CODES}

    days = roster.days
    n_days = len(days)
    first_day_col = 4   # columns 1-3 are Engineer/Designation/SME

    def styled_row(values, is_header):
        # Convert a table row into (value, style_id) pairs with the right styles.
        out = []
        for c, value in enumerate(values, 1):
            value = _clean(value)
            day_idx = c - first_day_col
            if is_header:
                if 0 <= day_idx < n_days:
                    out.append((value, header_styles[cfg.day_type(days[day_idx])]))
                else:
                    out.append((value, header_plain if value else name_style))
            elif 0 <= day_idx < n_days:
                # A grid cell: colour it by its code; otherwise plain.
                out.append((value, code_styles.get(value, num_style)))
            elif c >= first_day_col:
                # A summary column: centred number.
                out.append((int(value) if value.isdigit() else value, num_style))
            else:
                out.append((value, name_style))
        return out

    rows_src = table_rows(roster)
    sheet_rows = [styled_row(r, is_header=(i < 3)) for i, r in enumerate(rows_src)]

    # A legend a couple of rows below the table.
    sheet_rows.append([])
    sheet_rows.append([("Legend", header_plain)])
    for code in ALL_CODES:
        sheet_rows.append([(code, code_styles[code]), (CODE_DESCRIPTIONS[code], name_style)])

    # Column widths: wide name columns, each day column sized to its on-call names.
    widths = {1: 20, 2: 16, 3: 9}
    for i, d in enumerate(days):
        widths[first_day_col + i] = max(5, len(roster.primary.get(d, "")) + 1, len(roster.secondary.get(d, "")) + 1)

    sheets = [("Roster", _sheet_xml(sheet_rows, freeze=(3, 3), col_widths=widths))]

    # An optional second sheet listing the issues.
    if issues is not None:
        issue_rows = [[("Severity", header_plain), ("Date", header_plain), ("Message", header_plain)]]
        for i in issues:
            issue_rows.append([(i.severity, name_style), (i.day.isoformat() if i.day else "", name_style), (_clean(i.message), name_style)])
        if not issues:
            issue_rows.append([("OK", name_style), ("", name_style), ("All mandatory rules are satisfied", name_style)])
        sheets.append(("Issues", _sheet_xml(issue_rows, col_widths={3: 90})))

    # Assemble the package parts.
    content_types = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                     '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                     '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                     '<Default Extension="xml" ContentType="application/xml"/>'
                     '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
                     '<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>'
                     + "".join(f'<Override PartName="/xl/worksheets/sheet{i}.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
                               for i in range(1, len(sheets) + 1))
                     + '</Types>')
    root_rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                 '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                 '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>'
                 '</Relationships>')
    # Sheet names (Excel limits them to 31 chars and forbids some characters).
    def safe_name(name):
        return re.sub(r'[\\/*?:\[\]]', " ", name)[:31]

    sheet_names = [safe_name(f"{sheets[0][0]} {cfg.year}-{cfg.month:02d}")] + [safe_name(n) for n, _ in sheets[1:]]
    workbook = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
                'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets>'
                + "".join(f'<sheet name="{_xml_escape(nm)}" sheetId="{i}" r:id="rId{i}"/>' for i, nm in enumerate(sheet_names, 1))
                + '</sheets></workbook>')
    workbook_rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                     '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                     + "".join(f'<Relationship Id="rId{i}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet{i}.xml"/>'
                               for i in range(1, len(sheets) + 1))
                     + f'<Relationship Id="rId{len(sheets) + 1}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>'
                     + '</Relationships>')

    # Write everything into the .xlsx ZIP.
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", content_types)
        z.writestr("_rels/.rels", root_rels)
        z.writestr("xl/workbook.xml", workbook)
        z.writestr("xl/_rels/workbook.xml.rels", workbook_rels)
        z.writestr("xl/styles.xml", st.xml())
        for i, (_name, xml) in enumerate(sheets, 1):
            z.writestr(f"xl/worksheets/sheet{i}.xml", xml)

